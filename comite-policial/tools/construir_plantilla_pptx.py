"""Genera plantillas/plantilla_comite.pptx a partir de un informe PPT real.

Uso:
    python tools/construir_plantilla_pptx.py "COMITE POLICIAL 31.AGO AL 06.SEP.026.pptx"

Conserva el diseño exacto (fondo, tablas, cuadros de texto, logos y fotos
institucionales de la portada) y quita lo que cambia cada semana: las fotos de
las láminas, la tabla de la sección III-D (venía pegada como imagen) y las
láminas de diligencias, de las que queda una sola como prototipo.
Ejecutarlo solo si cambia el diseño del informe; la plantilla se versiona.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

RAIZ = Path(__file__).resolve().parents[1]
DESTINO = RAIZ / "plantillas" / "plantilla_comite.pptx"

# Láminas (1-based) que se conservan: 1-8 fijas, 13 = prototipo de diligencia
# y 16 = cierre "PDI".
PROTOTIPO, CIERRE = 13, 16
QUITAR_LAMINAS = [9, 10, 11, 12, 14, 15]


def quitar_lamina(prs, idx0: int):
    ids = prs.slides._sldIdLst
    sld = list(ids)[idx0]
    prs.part.drop_rel(sld.rId)
    ids.remove(sld)


def main(origen: str):
    prs = Presentation(origen)
    slides = list(prs.slides)
    assert len(slides) == 16, "se esperaba el informe de 16 láminas"

    # Fotos y tablas-imagen de la semana: fuera (la portada conserva las suyas)
    for n, sl in enumerate(slides, 1):
        if n == 1:
            continue
        for sh in list(sl.shapes):
            if sh.shape_type == 13:
                rid = sh._element.blip_rId
                sh._element.getparent().remove(sh._element)
                sl.part.drop_rel(rid)   # sin esto la imagen sigue empaquetada

    # Prototipo de diligencia: texto neutro (el exportador lo reescribe)
    proto = slides[PROTOTIPO - 1]
    for sh in proto.shapes:
        if sh.name == "CuadroTexto 5":
            sh.text_frame.paragraphs[0].runs[0].text = "UNIDAD"
            for r in sh.text_frame.paragraphs[1].runs[1:]:
                r._r.getparent().remove(r._r)
            sh.text_frame.paragraphs[1].runs[0].text = "Título"
        elif sh.name == "CuadroTexto 7":
            for p in list(sh.text_frame.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            p0 = sh.text_frame.paragraphs[0]
            for r in p0.runs[1:]:
                r._r.getparent().remove(r._r)
            p0.runs[0].text = "Fecha: "

    for n in sorted(QUITAR_LAMINAS, reverse=True):
        quitar_lamina(prs, n - 1)

    prs.core_properties.title = "Informe Comité Policial"
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    DESTINO.parent.mkdir(exist_ok=True)
    prs.save(DESTINO)
    print(f"{DESTINO}  ({DESTINO.stat().st_size / 1024:.0f} KB, "
          f"{len(prs.slides)} láminas)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
