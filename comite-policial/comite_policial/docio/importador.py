"""Lectura del informe semanal (.docx) hacia el modelo de datos de la aplicación."""
import copy
import io
import re
from docx import Document
from docx.oxml.ns import qn

from comite_policial.core.recuadros import (TIPO_IMPORTADO, contenido)
from comite_policial.docio.mapa import detectar, bloque_texto_x
from comite_policial.docio.semana import semana_normalizada
from comite_policial.docio.xml import (W_P, W_TR, cell_text, par_runs, par_text,
                                       phys_cells, norm)

ETIQUETA_RE = re.compile(r"^\s*([A-ZÁÉÍÓÚÑ][^:\n]{1,60}?)\s*:\s*(.*)$", re.DOTALL)
DETENIDOS_RE = re.compile(r"^(\d+)\s+DETENIDOS?\s*$", re.IGNORECASE)


def abrir(origen) -> Document:
    """`origen`: ruta o bytes."""
    if isinstance(origen, (bytes, bytearray)):
        return Document(io.BytesIO(origen))
    return Document(origen)


def _fila(tr):
    return [cell_text(c) for c in phys_cells(tr)]


def _parse_izquierda(tc):
    """(unidad, grupo, detenidos, titulo) a partir de la celda izquierda."""
    bloques, cur = [], []
    for p in tc.findall(W_P):
        for ln in par_text(p).split("\n"):
            if ln.strip():
                cur.append(" ".join(ln.split()))
            elif cur:
                bloques.append(cur)
                cur = []
    if cur:
        bloques.append(cur)
    unidad = " ".join(bloques[0]) if bloques else ""
    grupo, det, titulo = "", "", []
    for b in bloques[1:]:
        txt = " ".join(b)
        m = DETENIDOS_RE.match(txt)
        if m and not det:
            det = m.group(1)
        elif txt.upper().startswith("GRUPO") and len(b) == 1 and not grupo:
            grupo = txt
        else:
            titulo.extend(b)
    return unidad, grupo, det, "\n".join(titulo)


def _parse_derecha(tc):
    """(campos, defs) a partir de la celda derecha (etiquetas en negrita)."""
    campos, defs, actual = {}, [], None
    for p in tc.findall(W_P):
        texto = par_text(p).strip()
        if not texto:
            continue
        runs = par_runs(p)
        m = ETIQUETA_RE.match(texto)
        es_etiqueta = bool(m) and bool(runs) and runs[0][1]
        if es_etiqueta:
            etiqueta, resto = m.group(1).strip(), m.group(2).strip()
            base, k = etiqueta, 2
            while etiqueta in campos:
                etiqueta, k = f"{base} ({k})", k + 1
            clase = ("linea" if norm(base) in ("FECHA", "LUGAR")
                     else ("parrafo" if resto else "lista"))
            defs.append((etiqueta, clase))
            campos[etiqueta] = resto
            actual = etiqueta
        elif actual is not None:
            campos[actual] = (campos[actual] + "\n" + texto).strip("\n")
    return campos, defs


def parse_recuadro(tbl, indice: int) -> dict:
    tr = tbl.findall(W_TR)[0]
    izq, der = phys_cells(tr)
    unidad, grupo, det, titulo = _parse_izquierda(izq)
    campos, defs = _parse_derecha(der)
    rec = {"tipo": TIPO_IMPORTADO, "unidad": unidad, "grupo": grupo,
           "detenidos": det, "titulo": titulo, "campos": campos,
           "defs": defs, "_orig": indice}
    rec["_snap"] = copy.deepcopy(contenido(rec))
    return rec


def importar(origen) -> dict:
    """Lee el documento y devuelve el estado completo de la aplicación."""
    doc = abrir(origen)
    estado = {"semana": "", "turno": "", "tablas": {}, "recuadros": [],
              "texto_x": ""}

    for p in doc.paragraphs:
        t = p.text.strip()
        if not estado["semana"] and norm(t).startswith("SEMANA DEL"):
            estado["semana"] = semana_normalizada(t)
        elif not estado["turno"] and norm(t).startswith("TURNO"):
            estado["turno"] = t.split(":", 1)[1].strip() if ":" in t else t

    segs, recs = detectar(doc)
    for key, seg in segs.items():
        estado["tablas"][key] = {
            "rows": [_fila(tr) for tr in seg.data],
            "total": _fila(seg.total) if seg.total is not None else None,
            "fuente": (_fila(seg.fuente)[0] if seg.fuente is not None else ""),
            "ncell": seg.ncell,
        }
    estado["recuadros"] = [parse_recuadro(t, i) for i, t in enumerate(recs)]
    estado["texto_x"] = "\n".join(par_text(p).strip() for p in bloque_texto_x(doc))
    return estado
