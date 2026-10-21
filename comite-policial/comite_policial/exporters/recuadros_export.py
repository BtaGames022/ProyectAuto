"""Renderizado de los recuadros de la sección IX en Word y PowerPoint."""
from docx.shared import Pt, RGBColor, Inches as DInches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx.util import Inches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn as pqn
from lxml import etree

from comite_policial.core.recuadros import (
    lineas_planas_izquierda, parrafos_derecha)

TITULO_SECCION = "IX. CONCURRENCIAS Y DILIGENCIAS RELEVANTES"
FUENTE = "Arial"


# ── WORD ──────────────────────────────────────────────────────────────────────
def _bordes_word(table, sz=8):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)


def _run(par, texto, negrita=False, size=9.5):
    r = par.add_run(texto)
    r.bold = negrita
    r.font.name = FUENTE
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    r.font.size = Pt(size)
    return r


def crear_recuadro_word(doc, rec):
    """Agrega al final del documento un recuadro (tabla 1x2) y devuelve su elemento."""
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    _bordes_word(t)
    izq, der = t.rows[0].cells
    izq.width, der.width = DInches(1.95), DInches(4.55)

    # Columna izquierda: centrada, negrita, bloques separados por línea en blanco
    izq.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i, linea in enumerate(lineas_planas_izquierda(rec)):
        p = izq.paragraphs[0] if i == 0 else izq.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if linea:
            _run(p, linea, True, 10)

    # Columna derecha: etiquetas en negrita, texto justificado
    p = der.paragraphs[0]
    for i, par in enumerate(parrafos_derecha(rec)):
        if i:
            p = der.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(0)
        for texto, negrita in par:
            _run(p, texto, negrita)
    return t._tbl


def agregar_recuadros_word(doc, recuadros: list):
    if not recuadros:
        return
    h = doc.add_heading(TITULO_SECCION, level=1)
    h.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    for rec in recuadros:
        crear_recuadro_word(doc, rec)
        sep = doc.add_paragraph()
        sep.paragraph_format.space_after = Pt(6)


# ── POWERPOINT ────────────────────────────────────────────────────────────────
def _bordes_pptx(cell, color="000000", ancho=12700):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.SubElement(tcPr, pqn(tag), w=str(ancho), cap="flat",
                              cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, pqn("a:solidFill"))
        etree.SubElement(fill, pqn("a:srgbClr"), val=color)
        tcPr.remove(ln)
        tcPr.insert(i, ln)


def _tamano_fuente(rec) -> int:
    n = sum(len(t) for par in parrafos_derecha(rec) for t, _ in par)
    n_lineas = len(parrafos_derecha(rec))
    carga = n + n_lineas * 60
    if carga < 650:
        return 12
    if carga < 1100:
        return 10
    if carga < 1600:
        return 9
    return 8


def agregar_recuadros_pptx(prs, recuadros: list, slide_titulo, blank):
    """Una diapositiva por recuadro (tabla 1x2 con bordes negros)."""
    for rec in recuadros:
        sl = blank()
        slide_titulo(sl, TITULO_SECCION)
        fs = _tamano_fuente(rec)
        shape = sl.shapes.add_table(1, 2, Inches(0.2), Inches(0.8),
                                    Inches(9.6), Inches(4.5))
        tbl = shape.table
        tbl.columns[0].width = Inches(2.6)
        tbl.columns[1].width = Inches(7.0)
        # Quitar estilo/bandas por defecto
        tblPr = shape._element.graphic.graphicData.tbl.tblPr
        tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")

        izq, der = tbl.cell(0, 0), tbl.cell(0, 1)
        for c in (izq, der):
            c.fill.solid()
            c.fill.fore_color.rgb = PptRGB(0xFF, 0xFF, 0xFF)
            _bordes_pptx(c)
            c.margin_left = c.margin_right = Inches(0.08)

        # Izquierda
        izq.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = izq.text_frame
        tf.word_wrap = True
        for i, linea in enumerate(lineas_planas_izquierda(rec)):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = PP_ALIGN.CENTER
            r = par.add_run(); r.text = linea if linea else " "
            r.font.bold = True; r.font.size = PptPt(fs + 1)
            r.font.name = FUENTE; r.font.color.rgb = PptRGB(0, 0, 0)
        # Derecha
        der.vertical_anchor = MSO_ANCHOR.TOP
        tf = der.text_frame
        tf.word_wrap = True
        for i, parrafo in enumerate(parrafos_derecha(rec)):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = PP_ALIGN.JUSTIFY
            if not parrafo:
                r = par.add_run(); r.text = " "
                r.font.size = PptPt(max(fs - 4, 4))
                continue
            for texto, negrita in parrafo:
                r = par.add_run(); r.text = texto
                r.font.bold = negrita; r.font.size = PptPt(fs)
                r.font.name = FUENTE; r.font.color.rgb = PptRGB(0, 0, 0)
