"""Exportación a Microsoft PowerPoint (.pptx)."""
import os
import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN
from comite_policial.config.settings import settings
from comite_policial.exporters.recuadros_export import agregar_recuadros_pptx
from comite_policial.config.catalogos import *
from comite_policial.core.calculos import *


class PptxExporterMixin:
    # ── POWERPOINT ────────────────────────────────────────────────────────────
    def _gen_pptx(self, directorio: str) -> str:
        prs = Presentation()

        def blank(): return prs.slides.add_slide(prs.slide_layouts[6])

        def slide_titulo(slide, texto):
            tb = slide.shapes.add_textbox(
                Inches(0.2), Inches(0.06), Inches(9.6), Inches(0.58))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = texto; p.alignment = PP_ALIGN.LEFT
            r = p.runs[0]
            r.font.size = PptPt(16); r.font.bold = True
            r.font.color.rgb = PptRGB(0xFF, 0xFF, 0xFF)
            tb.fill.solid()
            tb.fill.fore_color.rgb = PptRGB(0x1E, 0x3A, 0x5F)

        def tabla_pptx(slide, filas, top, left, ancho, alto):
            if not filas: return
            nr, nc = len(filas), len(filas[0])
            tbl = slide.shapes.add_table(
                nr, nc, Inches(left), Inches(top), Inches(ancho), Inches(alto)).table
            for ri, row_data in enumerate(filas):
                is_total  = str(row_data[0]).startswith("TOTAL")
                is_fuente = str(row_data[0]).startswith("FUENTE")
                for ci, val in enumerate(row_data):
                    cell = tbl.cell(ri, ci)
                    cell.text = str(val)
                    tf2 = cell.text_frame.paragraphs[0]
                    tf2.alignment = PP_ALIGN.CENTER
                    run = tf2.runs[0] if tf2.runs else tf2.add_run()
                    run.font.size = PptPt(9)
                    if ri == 0:
                        run.font.bold = True
                        run.font.color.rgb = PptRGB(0xFF, 0xFF, 0xFF)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = PptRGB(0x1E, 0x3A, 0x5F)
                    elif is_total:
                        run.font.bold = True
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = PptRGB(0xDB, 0xEA, 0xFE)
                    elif is_fuente:
                        run.font.italic = True
                        run.font.size = PptPt(7)
                        run.font.color.rgb = PptRGB(0x64, 0x74, 0x8B)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = PptRGB(0xF8, 0xFA, 0xFC)

        def make_slide(titulo, filas, top=0.8, left=0.2,
                       ancho=9.6, alto=None):
            sl = blank()
            slide_titulo(sl, titulo)
            h = alto if alto else min(0.38 * len(filas) + 0.1, 6.5)
            tabla_pptx(sl, filas, top, left, ancho, h)
            return sl

        # Portada
        sl = blank()
        tb = sl.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(3))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = "COMITÉ POLICIAL SEMANAL"
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size = PptPt(36)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = PptRGB(0x1E, 0x3A, 0x5F)
        for line in [f"Región Policial Valparaíso",
                     f"Semana del {self.v_semana.get()}",
                     f"Turno: {self.v_turno.get()}"]:
            np_ = tf.add_paragraph()
            np_.text = line
            np_.alignment = PP_ALIGN.CENTER
            if np_.runs:
                np_.runs[0].font.size = PptPt(16)
                np_.runs[0].font.color.rgb = PptRGB(0x47, 0x55, 0x69)

        make_slide(f"I. ARMAS DE LA SEMANA — {self.v_semana.get()}",
                   self._get_tree_data(self.t_armas))
        make_slide(f"II. DROGA INCAUTADA — {self.v_semana.get()} (Gramos)",
                   self._get_tree_data(self.t_drogas_g))
        make_slide("II. DROGA — Plantas",
                   self._get_tree_data(self.t_drogas_p))
        make_slide(f"III‑A. HOMICIDIOS SEMANA — {self.v_semana.get()}",
                   self._get_tree_data(self.t_homi_a))
        make_slide("III‑B. MODALIDAD EN HOMICIDIOS",
                   self._get_tree_data(self.t_homi_b), ancho=9.6)
        make_slide("III‑C. RESULTADO INVESTIGATIVO (% base 100)",
                   self._get_tree_data(self.t_homi_c))
        make_slide(f"III‑D. COMPARATIVO {PERIODO} {ANO_ANT}/{ANO_ACT}",
                   self._get_tree_data(self.t_homi_d))
        make_slide("IV. MIGRACIONES Y POLICÍA INTERNACIONAL",
                   self._get_tree_data(self.t_migr), ancho=9.6)
        make_slide("IV‑B. FLUJO MIGRATORIO — LOS LIBERTADORES",
                   self._get_tree_data(self.t_flujo), ancho=9.6)
        make_slide("V. VEHÍCULOS — ESTADO",
                   self._get_tree_data(self.t_veh_e))
        make_slide("V. VEHÍCULOS — TIPO",
                   self._get_tree_data(self.t_veh_t))
        make_slide("VI. DELITOS SEXUALES (BRISEX)",
                   self._get_tree_data(self.t_brisex))
        make_slide(f"VII. BANDAS DESARTICULADAS — {self.v_semana.get()}",
                   self._get_tree_data(self.t_bandas))
        make_slide(f"VIII‑A. DETENIDOS POR DELITO — {self.v_semana.get()}",
                   self._get_tree_data(self.t_det_a))
        make_slide(f"VIII‑B. DETENIDOS POR PREFECTURA — {self.v_semana.get()}",
                   self._get_tree_data(self.t_det_b))
        make_slide(f"VIII‑C. DETENIDOS {ANO_ACT} A LA FECHA",
                   self._get_tree_data(self.t_det_c))

        agregar_recuadros_pptx(prs, getattr(self, "recuadros", []),
                               slide_titulo, blank)

        texto_x = self.txt_x.get("1.0", "end-1c").strip()
        if texto_x:
            sx = blank()
            slide_titulo(sx, "X. DILIGENCIAS DE RELEVANCIA DE LA SEMANA")
            tbx = sx.shapes.add_textbox(Inches(0.4), Inches(1.0), Inches(9.2), Inches(5))
            tbx.text_frame.word_wrap = True
            for i, linea in enumerate(texto_x.split("\n")):
                par = tbx.text_frame.paragraphs[0] if i == 0 else tbx.text_frame.add_paragraph()
                par.text = linea
                par.runs[0].font.size = PptPt(16)

        nombre = f"PRESENTACION_MANDO_{self.v_fecha.get().replace('/','‑')}.pptx"
        ruta = os.path.join(directorio, nombre)
        prs.save(ruta)
        return ruta
