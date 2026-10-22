"""Informe semanal en PowerPoint con el diseño exacto del comité.

No se dibuja nada desde cero: se abre `plantillas/plantilla_comite.pptx` (el
informe real sin las fotos de la semana) y se reescriben solo los textos y las
celdas, de modo que fondo, fuentes, colores, bordes y posiciones son los del
original. Las láminas de diligencias se clonan de una lámina prototipo, una
por recuadro de la sección IX.

Este módulo no depende de la interfaz: recibe el estado como diccionarios.
"""
from __future__ import annotations

import copy
import io
import math
import re

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageOps

from comite_policial.core.recuadros import campos_def

ALTO_LAMINA = 6858000
ANCHO_LAMINA = 12192000

# ── Espacios para fotos (EMU) ────────────────────────────────────────────────
# (clave, etiqueta, lámina, (x, y, ancho, alto)): son los marcos de las fotos
# del informe original; la foto elegida se recorta para llenar el marco.
SLOTS_FOTOS = [
    ("armas_1", "I. Armas — foto", 2, (7346196, 707942, 3670403, 4797507)),
    ("droga_1", "II. Droga — foto superior", 3, (6848475, 447853, 4580255, 2581097)),
    ("droga_2", "II. Droga — foto inferior", 3, (6848475, 3432909, 4580255, 2605942)),
    ("homi_1", "III. Homicidios — foto", 4, (7901683, 3551189, 4003496, 2468247)),
    ("migr_1", "IV. Migraciones — foto superior", 5, (8729978, 672842, 3003180, 2756158)),
    ("migr_2", "IV. Migraciones — foto inferior", 5, (8729977, 3587910, 3003179, 2756159)),
    ("flujo_1", "IV-B. Flujo — foto izquierda", 6, (1270009, 4674216, 2606666, 1737808)),
    ("flujo_2", "IV-B. Flujo — foto central", 6, (4693509, 4671266, 2700716, 1740759)),
    ("flujo_3", "IV-B. Flujo — foto derecha", 6, (8489491, 4741604, 2700716, 1740759)),
    ("vehban_1", "V-VII. Vehículos y bandas — foto izquierda", 7, (418527, 3314701, 2915432, 2727808)),
    ("vehban_2", "V-VII. Vehículos y bandas — foto derecha", 7, (3630883, 3314701, 2541317, 2727807)),
    ("det_1", "VIII. Detenidos — foto superior", 8, (6610350, 514048, 4908309, 2553002)),
    ("det_2", "VIII. Detenidos — foto inferior", 8, (6610349, 3324860, 4908309, 3088846)),
]

# Región de fotos de las láminas de diligencias
REGION_FOTOS = (7250000, 760000, 4450000, 5650000)
SEPARACION = 90000

# ── Qué tabla de la plantilla recibe cada tabla de la aplicación ─────────────
# clave de la app → (lámina, nombre del objeto, filas de cabecera, hay total,
#                    hay fuente, filas de texto libre que crecen)
TABLAS = {
    "armas":    (2, "Tabla 6", 1, True, True, False),
    "drogas_p": (3, "Tabla 13", 1, True, True, False),
    "drogas_g": (3, "Tabla 4", 1, True, True, False),
    "homi_a":   (4, "Tabla 11", 1, True, True, False),
    "homi_c":   (4, "Tabla 12", 1, True, True, False),
    "homi_b":   (4, "Tabla 14", 2, True, True, False),
    "migr":     (5, "Tabla 6", 1, True, False, True),
    "flujo":    (6, "Tabla 5", 1, True, False, False),
    "bandas":   (7, "Tabla 5", 1, False, True, True),
    "veh_e":    (7, "Tabla 2", 1, True, True, False),
    "veh_t":    (7, "Tabla 7", 1, True, True, False),
    "det_b":    (8, "Tabla 3", 1, True, True, False),
    "det_c":    (8, "Tabla 4", 1, True, True, False),
}


# ═════════════════════════════════════════════════════════════════════════════
# TEXTO
# ═════════════════════════════════════════════════════════════════════════════
def _set_par(p, texto: str):
    """Pone `texto` en un <a:p> conservando el formato de su primer run."""
    runs = p.findall(qn("a:r"))
    if runs:
        runs[0].find(qn("a:t")).text = texto
        for r in runs[1:]:
            p.remove(r)
        for br in p.findall(qn("a:br")):
            p.remove(br)
        return
    r = etree.Element(qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"))
    fin = p.find(qn("a:endParaRPr"))
    if fin is not None:
        for k, v in fin.attrib.items():
            rpr.set(k, v)
        for hijo in fin:
            rpr.append(copy.deepcopy(hijo))
    etree.SubElement(r, qn("a:t")).text = texto
    if fin is not None:
        fin.addprevious(r)
    else:
        p.append(r)


def _set_celda(tc, texto):
    """Reemplaza el contenido de una celda; cada línea es un párrafo."""
    cuerpo = tc.find(qn("a:txBody"))
    ps = cuerpo.findall(qn("a:p"))
    # molde: el párrafo con texto real (los vacíos/«\xa0» pueden tener otro color)
    def _txt(p):
        return "".join(t.text or "" for t in p.iter(qn("a:t"))).replace("\xa0", "").strip()
    base = next((p for p in ps if _txt(p)), None)
    if base is None:
        base = next((p for p in ps if p.find(qn("a:r")) is not None), ps[0])
    lineas = str(texto).split("\n") or [""]
    for p in ps:
        cuerpo.remove(p)
    for ln in lineas:
        p = copy.deepcopy(base)
        _set_par(p, ln)
        cuerpo.append(p)


def _tam_celda(tc) -> float:
    for rpr in tc.iter(qn("a:rPr")):
        if rpr.get("sz"):
            return int(rpr.get("sz")) / 100
    return 10.0


def _lineas(texto: str, pt: float, ancho_emu: int) -> int:
    """Líneas que ocupa `texto` a `pt` puntos en `ancho_emu` (estimación Arial)."""
    util = max(ancho_emu / 12700 - 14.4, 20)
    por_linea = max(util / (0.53 * pt), 4)
    total = 0
    for parrafo in str(texto).split("\n"):
        n, actual = 1, 0
        for palabra in parrafo.split(" "):
            largo = len(palabra) + (1 if actual else 0)
            if actual and actual + largo > por_linea:
                n += 1
                actual = len(palabra)
            else:
                actual += largo
        total += n
    return total


# ═════════════════════════════════════════════════════════════════════════════
# TABLAS
# ═════════════════════════════════════════════════════════════════════════════
def _fila_clon(tr):
    nueva = copy.deepcopy(tr)
    for ext in nueva.findall(qn("a:extLst")):
        nueva.remove(ext)
    return nueva


def _llenar_tabla(shape, datos, total, *, n_cab=1, hay_total=True,
                  hay_fuente=True, cab=None, texto_libre=False,
                  presupuesto=None):
    """Vuelca `datos` (+ `total`) en la tabla conservando el estilo de cada fila.

    Agrega o quita filas de datos clonando la última fila de datos.
    `cab`: {(fila, col): texto} para encabezados que llevan la semana o el año.
    """
    tbl = shape._element.graphic.graphicData.tbl
    trs = tbl.findall(qn("a:tr"))
    ini = n_cab
    fin = len(trs) - (1 if hay_total else 0) - (1 if hay_fuente else 0)
    n_plantilla, n_datos = fin - ini, len(datos)

    if n_datos > n_plantilla and n_plantilla > 0:
        ref = trs[fin - 1]
        for _ in range(n_datos - n_plantilla):
            nueva = _fila_clon(ref)
            ref.addnext(nueva)
            ref = nueva
    elif n_datos < n_plantilla:
        for tr in trs[ini + n_datos:fin]:
            tbl.remove(tr)
    trs = tbl.findall(qn("a:tr"))

    ncols = len(tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol")))

    def escribir(tr, valores, es_total=False):
        for ci, tc in enumerate(tr.findall(qn("a:tc"))):
            if tc.get("hMerge") or tc.get("vMerge") or ci >= ncols:
                continue
            if es_total and ci == 0:
                continue                  # el rótulo «TOTAL …» es el de la plantilla
            if ci < len(valores):
                _set_celda(tc, valores[ci])
                if texto_libre:
                    tc.find(qn("a:tcPr")).set("anchor", "ctr")
                    if es_total:          # la plantilla centraba con espacios
                        for pr in tc.iter(qn("a:pPr")):
                            pr.set("algn", "ctr")

    for i, fila in enumerate(datos):
        escribir(trs[ini + i], list(fila))
    if hay_total and total is not None:
        escribir(trs[ini + n_datos], list(total), es_total=True)
    for (r, c), txt in (cab or {}).items():
        _set_celda(trs[r].findall(qn("a:tc"))[c], txt)

    _ajustar_altos(shape, tbl, n_cab, n_datos, hay_total, hay_fuente,
                   texto_libre, presupuesto)


def _ajustar_altos(shape, tbl, n_cab, n_datos, hay_total, hay_fuente,
                   texto_libre, presupuesto):
    """Altura de filas: crece con el texto libre o se comprime al presupuesto."""
    trs = tbl.findall(qn("a:tr"))
    anchos = [int(g.get("w")) for g in tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))]
    ini, fin = n_cab, n_cab + n_datos

    if texto_libre:
        for tr in trs[ini:]:
            necesario = 0
            for ci, tc in enumerate(tr.findall(qn("a:tc"))):
                if tc.get("hMerge") or tc.get("vMerge"):
                    continue
                span = int(tc.get("gridSpan", "1"))
                ancho = sum(anchos[ci:ci + span])
                pt = _tam_celda(tc)
                txt = "\n".join("".join(t.text or "" for t in p.iter(qn("a:t")))
                                for p in tc.iter(qn("a:p")))
                necesario = max(necesario, _lineas(txt, pt, ancho) * pt * 1.2 * 12700 + 60000)
            tr.set("h", str(int(max(necesario, 200000))))

    if presupuesto:
        altos = [int(tr.get("h")) for tr in trs]
        exceso = sum(altos) - presupuesto
        if exceso > 0 and n_datos:
            # primero la cabecera, luego las filas de datos (mínimo para 10-12 pt)
            cab_min = 560000
            for i in range(n_cab):
                rec = min(altos[i] - cab_min, exceso) if altos[i] > cab_min else 0
                altos[i] -= rec
                exceso -= rec
            if exceso > 0:
                piso = 280000
                for i in range(ini, fin):
                    rec = min(altos[i] - piso, math.ceil(exceso / (fin - i)))
                    rec = max(rec, 0)
                    altos[i] -= rec
                    exceso -= rec
            for tr, h in zip(trs, altos):
                tr.set("h", str(int(h)))

    shape.height = sum(int(tr.get("h")) for tr in trs)


def _tabla(slide, nombre: str):
    return next(sh for sh in slide.shapes if sh.has_table and sh.name == nombre)


# ═════════════════════════════════════════════════════════════════════════════
# FOTOS
# ═════════════════════════════════════════════════════════════════════════════
def _imagen_flujo(ruta: str):
    """Abre la foto respetando su rotación, la achica y devuelve (bytes, ancho, alto)."""
    with Image.open(ruta) as im:
        im = ImageOps.exif_transpose(im)
        im.thumbnail((1920, 1920))
        buf = io.BytesIO()
        if im.mode in ("RGBA", "LA", "P"):
            im.convert("RGBA").save(buf, "PNG")
        else:
            im.convert("RGB").save(buf, "JPEG", quality=88)
        buf.seek(0)
        return buf, im.size[0], im.size[1]


def _poner_foto(slide, ruta: str, x, y, w, h, modo="cubrir"):
    buf, iw, ih = _imagen_flujo(ruta)
    if modo == "cubrir":
        pic = slide.shapes.add_picture(buf, x, y, w, h)
        r_marco, r_img = w / h, iw / ih
        if r_img > r_marco:
            corte = (1 - r_marco / r_img) / 2
            pic.crop_left = pic.crop_right = corte
        else:
            corte = (1 - r_img / r_marco) / 2
            pic.crop_top = pic.crop_bottom = corte
    else:  # ajustar: sin recortar, centrada en el marco
        esc = min(w / iw, h / ih)
        w2, h2 = int(iw * esc), int(ih * esc)
        slide.shapes.add_picture(buf, x + (w - w2) // 2, y + (h - h2) // 2, w2, h2)


def _celdas_fotos(n: int, region=REGION_FOTOS):
    """Reparte la región de fotos en `n` celdas (x, y, w, h)."""
    x0, y0, W, H = region
    if n <= 0:
        return []
    if n == 1:
        filas = [[1]]
    elif n == 2:
        filas = [[1], [1]]
    elif n == 3:
        filas = [[1], [1, 1]]
    else:
        filas = [[1, 1] for _ in range(math.ceil(n / 2))]
        if n % 2:
            filas[-1] = [1]
    alto = (H - SEPARACION * (len(filas) - 1)) / len(filas)
    celdas, y = [], y0
    for fila in filas:
        ancho = (W - SEPARACION * (len(fila) - 1)) / len(fila)
        for k in range(len(fila)):
            celdas.append((int(x0 + k * (ancho + SEPARACION)), int(y),
                           int(ancho), int(alto)))
        y += alto + SEPARACION
    return celdas[:n]


# ═════════════════════════════════════════════════════════════════════════════
# LÁMINAS
# ═════════════════════════════════════════════════════════════════════════════
def _clonar_lamina(prs, origen):
    """Duplica una lámina (fondo incluido) al final de la presentación."""
    nueva = prs.slides.add_slide(origen.slide_layout)
    arbol = nueva.shapes._spTree
    for el in list(arbol)[2:]:                 # fuera marcadores del layout
        arbol.remove(el)
    mapa = {}
    for rid, rel in origen.part.rels.items():
        if rel.reltype == RT.IMAGE:
            mapa[rid] = nueva.part.relate_to(rel.target_part, rel.reltype)
    bg = origen._element.cSld.find(qn("p:bg"))
    if bg is not None:
        nbg = copy.deepcopy(bg)
        for el in nbg.iter():
            for k in list(el.attrib):
                if k.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"):
                    el.set(k, mapa[el.get(k)])
        nueva._element.cSld.insert(0, nbg)
    for el in list(origen.shapes._spTree)[2:]:
        arbol.append(copy.deepcopy(el))
    return nueva


def _quitar_lamina(prs, lamina):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        if prs.slides.get(int(sld.get("id"))) is lamina:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
            return


def _mover_al_final(prs, lamina):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        if prs.slides.get(int(sld.get("id"))) is lamina:
            ids.remove(sld)
            ids.append(sld)
            return


def _forma(slide, nombre, empieza=""):
    return next(sh for sh in slide.shapes
                if sh.name == nombre and sh.has_text_frame
                and sh.text_frame.text.startswith(empieza))


def _titulo_semana(slide, nombre, idx_par, semana, empieza=""):
    sh = _forma(slide, nombre, empieza)
    ps = sh._element.txBody.findall(qn("a:p"))
    _set_par(ps[idx_par], f"Semana del {semana}")


# ── Diligencias ──────────────────────────────────────────────────────────────
def _parrafos_diligencia(rec) -> list:
    """Campos del recuadro como párrafos [(texto, negrita)], con una línea en
    blanco entre campos (así están las láminas originales)."""
    campos = rec.get("campos", {})
    out = []
    for etiqueta, clase in campos_def(rec):
        valor = str(campos.get(etiqueta, "")).strip()
        if not valor:
            continue
        if out:
            out.append([])
        lineas = [ln.rstrip() for ln in valor.splitlines() if ln.strip()]
        if clase == "lista":
            out.append([(f"{etiqueta}:", True)])
            out.extend([(ln, False)] for ln in lineas)
        else:
            out.append([(f"{etiqueta}: ", True), (lineas[0], False)])
            out.extend([(ln, False)] for ln in lineas[1:])
    return out


def _run(texto, negrita, pt, color="260000"):
    r = etree.Element(qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"), lang="es-ES", sz=str(int(pt * 100)),
                           b="1" if negrita else "0", dirty="0")
    relleno = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(relleno, qn("a:srgbClr"), val=color)
    etree.SubElement(rpr, qn("a:latin"), typeface="Arial")
    etree.SubElement(rpr, qn("a:cs"), typeface="Arial")
    etree.SubElement(r, qn("a:t")).text = texto
    return r


def _reconstruir_parrafos(cuerpo, plantilla_ppr, parrafos, pt):
    """parrafos: lista de listas [(texto, negrita)]; [] = línea en blanco."""
    for p in cuerpo.findall(qn("a:p")):
        cuerpo.remove(p)
    for runs in parrafos:
        p = etree.SubElement(cuerpo, qn("a:p"))
        if plantilla_ppr is not None:
            p.append(copy.deepcopy(plantilla_ppr))
        for texto, negrita in runs:
            p.append(_run(texto, negrita, pt))
        fin = etree.SubElement(p, qn("a:endParaRPr"), lang="es-ES",
                               sz=str(int(pt * 100)), dirty="0")
        if not runs:
            etree.SubElement(etree.SubElement(fin, qn("a:solidFill")),
                             qn("a:srgbClr"), val="260000")


def _alto_parrafos(parrafos, pt, ancho_emu, espacio_despues_pt=0.0):
    n = 0
    for runs in parrafos:
        n += _lineas("".join(t for t, _ in runs), pt, ancho_emu) if runs else 1
    return int(n * pt * 1.2 * 12700 + len(parrafos) * espacio_despues_pt * 12700 + 91440)


def _lamina_diligencia(prs, proto, rec, avisos):
    sl = _clonar_lamina(prs, proto)
    enc = next(sh for sh in sl.shapes if sh.name == "CuadroTexto 5")
    cue = next(sh for sh in sl.shapes if sh.name == "CuadroTexto 7")
    fotos = [f for f in rec.get("fotos", []) if f]
    x_txt = 286802
    # con fotos el texto va en la columna izquierda; sin fotos puede ensancharse
    # si el relato es largo (se prefiere letra de 11 pt o más)
    anchos = ([REGION_FOTOS[0] - 250000 - x_txt] if fotos
              else [7600000, 9400000, 11400000])

    unidad = rec.get("unidad", "").strip().upper()
    if rec.get("grupo", "").strip():
        unidad = f"{unidad} {rec['grupo'].strip().upper()}".strip()
    titulo = [ln.strip() for ln in rec.get("titulo", "").splitlines() if ln.strip()]
    pars_enc = ([[(unidad, True)]] if unidad else []) + [[(t, True)] for t in titulo]
    pars_enc = pars_enc or [[(" ", True)]]
    pars = _parrafos_diligencia(rec) or [[("Sin datos.", False)]]

    for ancho in anchos:
        # Encabezado: unidad (+ grupo) y título del procedimiento
        for pt_enc in (14, 13, 12, 11):
            alto_enc = _alto_parrafos(pars_enc, pt_enc, ancho, 10.0)
            if alto_enc <= 1250000:
                break
        y_cuerpo = 713870 + alto_enc + 60000
        libre = ALTO_LAMINA - 230000 - y_cuerpo
        for pt_cue in (14, 13, 12, 11, 10, 9, 8):
            alto = _alto_parrafos(pars, pt_cue, ancho)
            if alto <= libre:
                break
        if alto <= libre and pt_cue >= 11:
            break
    if alto > libre:
        avisos.append(f"Diligencia «{unidad[:40]}»: el texto no cabe en una lámina; "
                      "conviene acortarlo.")

    tb = enc._element.txBody
    ppr_enc = tb.find(qn("a:p")).find(qn("a:pPr"))
    _reconstruir_parrafos(tb, ppr_enc, pars_enc, pt_enc)
    enc.left, enc.top, enc.width, enc.height = x_txt, 713870, ancho, alto_enc

    tb = cue._element.txBody
    ppr = tb.find(qn("a:p")).find(qn("a:pPr"))
    _reconstruir_parrafos(tb, ppr, pars, pt_cue)
    cue.left, cue.top, cue.width, cue.height = x_txt, y_cuerpo, ancho, min(alto, libre)

    for (x, y, w, h), ruta in zip(_celdas_fotos(len(fotos)), fotos):
        _poner_foto(sl, ruta, x, y, w, h, modo="ajustar")
    return sl


# ── III-D: la tabla venía pegada como imagen; se rehace como tabla nativa ────
def _tabla_comparativa(sl4, filas, total, cab):
    modelo = next(sh for sh in sl4.shapes if sh.name == "Tabla 12")
    el = copy.deepcopy(modelo._element)
    for ext in list(el.iter(qn("a:extLst"))):      # ids únicos: PowerPoint los regenera
        ext.getparent().remove(ext)
    modelo._element.addnext(el)
    ids = [int(i) for i in sl4._element.xpath("//p:cNvPr/@id")]
    cnv = el.xpath(".//p:cNvPr")[0]
    cnv.set("id", str(max(ids) + 1))
    cnv.set("name", "Tabla III-D")
    tbl = el.graphic.graphicData.tbl

    cols = tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))
    nueva = copy.deepcopy(cols[1])
    for e in nueva.findall(qn("a:extLst")):
        nueva.remove(e)
    cols[1].addnext(nueva)
    for tr in tbl.findall(qn("a:tr")):
        tcs = tr.findall(qn("a:tc"))
        tcs[1].addnext(copy.deepcopy(tcs[1]))
    anchos = [1250000, 780000, 780000, 734950]
    for g, w in zip(tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol")), anchos):
        g.set("w", str(w))

    ultimo = tbl.findall(qn("a:tr"))[-1].findall(qn("a:tc"))   # fila FUENTE
    ultimo[0].set("gridSpan", "4")
    for tc in ultimo[1:]:
        tc.set("hMerge", "1")

    forma = next(sh for sh in sl4.shapes if sh.name == "Tabla III-D")
    forma.left, forma.top, forma.width = 8360229, 1166332, sum(anchos)

    datos = [[f[0], f[1], f[2], ""] for f in filas]       # la variación va solo en el total
    _llenar_tabla(forma, datos, total, cab=cab)
    trs = tbl.findall(qn("a:tr"))
    if len(datos) > 1:                                    # columna VAR. fusionada
        trs[1].findall(qn("a:tc"))[3].set("rowSpan", str(len(datos)))
        for tr in trs[2:1 + len(datos)]:
            tr.findall(qn("a:tc"))[3].set("vMerge", "1")


# ═════════════════════════════════════════════════════════════════════════════
# API
# ═════════════════════════════════════════════════════════════════════════════
def exportar_informe(estado: dict, ruta_salida: str, plantilla: str,
                     fotos: dict | None = None, ano_ant="2025", ano_act="2026",
                     periodo="ENE - AGO") -> list:
    """Genera el PPTX. Devuelve la lista de avisos (textos que no caben, etc.).

    estado = {"semana": "31.AGO AL 06.SEP.026",
              "tablas": {clave: {"rows": [[...]], "total": [...] | None}},
              "recuadros": [dict, ...]}
    fotos  = {clave de SLOTS_FOTOS: ruta de la imagen}
    """
    avisos: list = []
    fotos = fotos or {}
    semana = estado["semana"].strip()
    fin_semana = semana.split(" AL ")[-1] if " AL " in semana else semana
    tablas = estado["tablas"]

    prs = Presentation(plantilla)
    lam = list(prs.slides)
    portada, resto = lam[0], lam[1:]
    s = {i + 2: sl for i, sl in enumerate(resto)}        # número de lámina → lámina
    proto, cierre = lam[8], lam[9]

    # Títulos con la semana
    _titulo_semana(portada, "Google Shape;125;p1", 2, semana)
    _titulo_semana(s[2], "Google Shape;132;p6", 1, semana)
    _titulo_semana(s[3], "Google Shape;142;p3", 1, semana)
    _titulo_semana(s[4], "Google Shape;154;p27", 1, semana)
    _titulo_semana(s[5], "Google Shape;167;p28", 2, semana, "MIGRACIONES")
    _titulo_semana(s[7], "Google Shape;179;p29", 1, semana, "BANDAS")
    _titulo_semana(s[7], "Google Shape;179;p29", 1, semana, "VEH")
    _titulo_semana(s[8], "Google Shape;179;p29", 1, semana, "DETENIDOS")

    a_fecha = f"{ano_act} A LA FECHA"
    sem_cab = f"SEMANA {semana}"
    cabeceras = {
        "armas":    {(0, 1): sem_cab, (0, 2): a_fecha},
        "drogas_p": {(0, 1): sem_cab, (0, 2): a_fecha},
        "drogas_g": {(0, 1): sem_cab, (0, 2): a_fecha},
        "homi_a":   {(0, 1): f"SEMANA DEL {semana}", (0, 2): f"AÑO {ano_act} A LA FECHA"},
        "homi_c":   {(0, 1): f"01.ENE AL {fin_semana}", (0, 2): f"% AÑO {ano_act}"},
        "homi_b":   {(0, 1): f"SEMANA DEL {semana}", (0, 4): f"AÑO {ano_act} A LA FECHA"},
        "veh_e":    {(0, 1): sem_cab, (0, 2): a_fecha},
        "veh_t":    {(0, 1): sem_cab, (0, 2): a_fecha},
        "det_b":    {(0, 3): sem_cab},
        "det_c":    {(0, 3): f"AÑO {ano_act} A LA FECHA"},
    }
    presupuestos = {"flujo": 2733048}

    for clave, (n_lam, nombre, n_cab, hay_total, hay_fuente, libre) in TABLAS.items():
        datos_t = tablas.get(clave)
        if not datos_t:
            continue
        _llenar_tabla(_tabla(s[n_lam], nombre), datos_t["rows"], datos_t.get("total"),
                      n_cab=n_cab, hay_total=hay_total, hay_fuente=hay_fuente,
                      cab=cabeceras.get(clave), texto_libre=libre,
                      presupuesto=presupuestos.get(clave))

    # III-D comparativo (tabla nativa en lugar de imagen)
    d = tablas.get("homi_d")
    if d:
        _tabla_comparativa(
            s[4], d["rows"], d.get("total"),
            {(0, 0): "HOMICIDIOS", (0, 1): f"{periodo} {ano_ant}",
             (0, 2): f"{periodo} {ano_act}", (0, 3): f"VAR. % {ano_ant} / {ano_act}"})

    # Avisos por desbordes de tablas
    for clave, (n_lam, nombre, *_r) in TABLAS.items():
        forma = _tabla(s[n_lam], nombre)
        if forma.top + forma.height > ALTO_LAMINA - 100000:
            avisos.append(f"La tabla «{clave}» se sale de la lámina {n_lam}; "
                          "reduce el texto de sus celdas.")

    # Fotos de las láminas fijas
    for clave, etiqueta, n_lam, (x, y, w, h) in SLOTS_FOTOS:
        ruta = fotos.get(clave)
        if ruta:
            try:
                _poner_foto(s[n_lam], ruta, x, y, w, h, modo="cubrir")
            except Exception as exc:
                avisos.append(f"No se pudo insertar la foto «{etiqueta}»: {exc}")

    # Láminas de diligencias (una por recuadro)
    for rec in estado.get("recuadros", []):
        _lamina_diligencia(prs, proto, rec, avisos)
    _quitar_lamina(prs, proto)
    _mover_al_final(prs, cierre)

    prs.core_properties.title = f"Informe Comité Policial — {semana}"
    prs.save(ruta_salida)
    return avisos
