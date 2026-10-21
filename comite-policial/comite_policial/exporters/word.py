"""Exportación a Microsoft Word (.docx)."""
import os
import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN
from comite_policial.config.settings import settings
from comite_policial.exporters.recuadros_export import agregar_recuadros_word
from comite_policial.config.catalogos import *
from comite_policial.core.calculos import *


class WordExporterMixin:
    # ── WORD ──────────────────────────────────────────────────────────────────
    def _gen_word(self, directorio: str) -> str:
        plantilla = settings.plantilla_docx
        doc = Document(plantilla) if os.path.exists(plantilla) else Document()

        def set_cell_bg(cell, hex_color):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), hex_color)
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:val"), "clear")
            tcPr.append(shd)

        def tabla_word(titulo, filas):
            if titulo:
                p = doc.add_heading(titulo, level=1)
                p.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
            if not filas:
                return
            ncols = len(filas[0])
            t = doc.add_table(rows=len(filas), cols=ncols)
            t.style = "Table Grid"
            for ri, row_data in enumerate(filas):
                cells = t.rows[ri].cells
                is_total  = str(row_data[0]).startswith("TOTAL")
                is_fuente = str(row_data[0]).startswith("FUENTE")
                for ci, val in enumerate(row_data):
                    cells[ci].text = str(val)
                    para = cells[ci].paragraphs[0]
                    para.alignment = (WD_ALIGN_PARAGRAPH.LEFT if ci == 0
                                      else WD_ALIGN_PARAGRAPH.CENTER)
                    if cells[ci].paragraphs[0].runs:
                        run = cells[ci].paragraphs[0].runs[0]
                        if ri == 0:  # Cabecera
                            run.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.font.size = Pt(10)
                            set_cell_bg(cells[ci], "1e3a5f")
                        elif is_total:
                            run.bold = True
                            set_cell_bg(cells[ci], "dbeafe")
                        elif is_fuente:
                            run.italic = True
                            run.font.size = Pt(8)
                            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
                            set_cell_bg(cells[ci], "f8fafc")
            doc.add_paragraph()

        # Encabezado
        h = doc.add_heading("INFORME SEMANAL COMITÉ POLICIAL.", 0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(
            f"Semana del {self.v_semana.get()}\n"
            f"Turno: {self.v_turno.get()}"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER

        tabla_word("I.- ARMAS DE LA SEMANA",
                   self._get_tree_data(self.t_armas))
        tabla_word("II.- DROGA INCAUTADA EN LA SEMANA (GRAMOS)",
                   self._get_tree_data(self.t_drogas_g))
        tabla_word("",
                   self._get_tree_data(self.t_drogas_p))
        tabla_word(f"III-A. HOMICIDIOS — SEMANA DEL {self.v_semana.get()}",
                   self._get_tree_data(self.t_homi_a))
        tabla_word(f"III-B. MODALIDAD EN HOMICIDIOS",
                   self._get_tree_data(self.t_homi_b))
        tabla_word("III-C. RESULTADO EN HOMICIDIOS (% INVESTIGATIVO)",
                   self._get_tree_data(self.t_homi_c))
        tabla_word(f"III-D. COMPARATIVO {PERIODO} {ANO_ANT}/{ANO_ACT}",
                   self._get_tree_data(self.t_homi_d))
        tabla_word("IV.- MIGRACIONES Y POLICÍA INTERNACIONAL",
                   self._get_tree_data(self.t_migr))
        tabla_word("IV-B. FLUJO MIGRATORIO COMPLEJO FRONTERIZO LOS LIBERTADORES",
                   self._get_tree_data(self.t_flujo))
        tabla_word("V.- VEHÍCULOS — POR ESTADO",
                   self._get_tree_data(self.t_veh_e))
        tabla_word("",
                   self._get_tree_data(self.t_veh_t))
        tabla_word("VI.- DELITOS SEXUALES (BRISEX)",
                   self._get_tree_data(self.t_brisex))
        tabla_word(f"VII.- BANDAS DESARTICULADAS — {self.v_semana.get()}",
                   self._get_tree_data(self.t_bandas))
        tabla_word(f"VIII-A. DETENIDOS Y ARRESTADOS — POR DELITO",
                   self._get_tree_data(self.t_det_a))
        tabla_word("VIII-B. DETENIDOS — POR PREFECTURA (SEMANA)",
                   self._get_tree_data(self.t_det_b))
        tabla_word(f"VIII-C. DETENIDOS — ACUMULADO {ANO_ACT} A LA FECHA",
                   self._get_tree_data(self.t_det_c))

        agregar_recuadros_word(doc, getattr(self, "recuadros", []))

        texto_x = self.txt_x.get("1.0", "end-1c").strip()
        if texto_x:
            hx = doc.add_heading("X. DILIGENCIAS DE RELEVANCIA DE LA SEMANA", level=1)
            hx.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
            for linea in texto_x.split("\n"):
                doc.add_paragraph(linea)

        nombre = f"INFORME_SEMANAL_{self.v_fecha.get().replace('/','‑')}.docx"
        ruta = os.path.join(directorio, nombre)
        doc.save(ruta)
        return ruta

