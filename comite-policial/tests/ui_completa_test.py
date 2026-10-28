"""Recorrido completo de la interfaz: cada vista, cada botón y la exportación.

Cualquier excepción dentro de un callback de Tk se registra y hace fallar la prueba
(por defecto Tk solo la imprime y la aplicación sigue «rota» sin avisar).
"""
import os
import sys
import tempfile
import zipfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
os.chdir(ROOT)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(__file__))
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from tkinter import filedialog, messagebox
from PIL import Image
from pptx import Presentation

import make_informe_docx
from comite_policial.ui.app import App
from comite_policial.ui.calendarios import CalendarPopup, WeekRangePopup
from comite_policial.ui.recuadros import CASILLA_NO, CASILLA_SI
from comite_policial.docio.importador import importar

errores = []
cuadros = []          # mensajes que la app mostraría al usuario
messagebox.showinfo = lambda *a, **k: cuadros.append(("info", a))
messagebox.showwarning = lambda *a, **k: cuadros.append(("warn", a))
messagebox.showerror = lambda *a, **k: cuadros.append(("error", a))
messagebox.askyesno = lambda *a, **k: False

tmp = tempfile.mkdtemp()
a = App()
a.geometry("1100x640")
a.report_callback_exception = lambda e, v, t: errores.append((e.__name__, str(v)))
a.update()


def pulsa(widget_boton):
    """Simula el clic sobre un Boton."""
    widget_boton._cmd()
    a.update()


def botones(w, out):
    from comite_policial.ui.widgets import Boton
    for c in w.winfo_children():
        if isinstance(c, Boton):
            out.append(c)
        botones(c, out)
    return out


def boton(vista, texto):
    for b in botones(a._vistas[vista], []):
        if b._lbl.cget("text") == texto:
            return b
    raise AssertionError(f"no hay botón «{texto}» en {vista}")


# ── 1. Todas las vistas se abren y la ventana pequeña no rompe nada ──────────
for clave in list(a._vistas):
    a._switch(clave)
    a.update()
assert not errores, errores
assert len(a._vistas) == 16, len(a._vistas)

# ── 2. Edición en la celda con teclado (Enter baja, Tab avanza) ──────────────
a._switch("armas")
t = a.t_armas
f0 = t.get_children()[0]
t._abrir_editor(f0, 1)
a.update()
ed = t._entry_popup.winfo_children()[0]
ed.delete(0, "end")
ed.insert(0, "7")
ed.event_generate("<Return>")
a.update()
assert t.raw_values(f0)[1] == "7", t.raw_values(f0)
total = [t.raw_values(r) for r in t.get_children() if "total" in t.item(r, "tags")][0]
assert total[1] == "7", total          # el total se recalcula
assert t._entry_popup is not None      # Enter abrió la celda de abajo
t._entry_popup.winfo_children()[0].event_generate("<Escape>")
a.update()
assert t._entry_popup is None

# ── 3. Flujo migratorio: agregar y eliminar días ─────────────────────────────
a._switch("flujo")
pulsa(boton("flujo", "Agregar día"))
pulsa(boton("flujo", "Agregar día"))
filas = [r for r in a.t_flujo.get_children()
         if not {"total", "fuente"} & set(a.t_flujo.item(r, "tags"))]
assert len(filas) == 2, len(filas)
a.t_flujo.selection_set(filas[0])
pulsa(boton("flujo", "Eliminar seleccionado"))
filas = [r for r in a.t_flujo.get_children()
         if not {"total", "fuente"} & set(a.t_flujo.item(r, "tags"))]
assert len(filas) == 1, len(filas)

# ── 4. Detenidos por delito: validación y valores no numéricos ───────────────
a._switch("det_a")
pulsa(boton("det_a", "Agregar fila"))                  # sin delito: avisa y no agrega
assert any(k == "warn" for k, _ in cuadros)
a._v_del_d.set("Ley 20.000")
a._v_del_a.set("abc")                                  # texto inválido
a._v_del_dt.set("5")
pulsa(boton("det_a", "Agregar fila"))
filas = [a.t_det_a.raw_values(r) for r in a.t_det_a.get_children()
         if not {"total", "fuente"} & set(a.t_det_a.item(r, "tags"))]
assert filas == [["Ley 20.000", "0", "5", "5"]], filas

# ── 5. Calendarios ───────────────────────────────────────────────────────────
import datetime
w = WeekRangePopup(a, a.v_semana)
w.update()
w._ok()                                                # sin selección: avisa
assert w.winfo_exists()
w._semana(-1)
w._ok()
a.update()
partes = a.v_semana.get().split(" AL ")
assert len(partes) == 2 and partes[1].count(".") == 2, a.v_semana.get()
lunes = datetime.date.today() - datetime.timedelta(days=datetime.date.today().weekday() + 7)
assert partes[0].startswith(f"{lunes.day:02d}."), (partes, lunes)
hoy = datetime.date.today()
c = CalendarPopup(a, a.v_fecha, "Fecha")
c._pick(datetime.date(2026, 9, 6))
c._ok()
assert a.v_fecha.get() == "06/09/2026", a.v_fecha.get()
a.v_semana.set("31.AGO AL 06.SEP.026")

# ── 6. Recuadros: crear, duplicar, mover, casilla PPT, eliminar ──────────────
a._switch("recuadros")
a.update()
a.rv_unidad.set("Brigada de Homicidios San Antonio")
a.rv_titulo.insert("1.0", "Homicidio con arma cortante")
a._rec_widgets["Lugar"][0].insert(0, "San Antonio")
a._rec_widgets["Diligencias"][0].insert("1.0", "Se recopilaron antecedentes.")
foto = os.path.join(tmp, "foto.jpg")
Image.new("RGB", (1600, 1000), (30, 60, 140)).save(foto)
a._rec_fotos = [foto]
a._rec_fotos_refrescar()
a._rec_guardar()
a.rv_unidad.set("")
a._rec_guardar()                                       # sin unidad: avisa y no guarda
assert len(a.recuadros) == 1
a.t_rec.selection_set(a.t_rec.get_children()[0])
a._rec_duplicar()
assert len(a.recuadros) == 2
a.recuadros[1]["unidad"] = "Brigada Antinarcóticos La Calera"
a._rec_refrescar_lista(1)
a.t_rec.selection_set(a.t_rec.get_children()[1])
a._rec_mover(-1)
assert a.recuadros[0]["unidad"].startswith("Brigada Antinarc")
assert a.recuadros[1].get("fotos") == [foto]           # las fotos viajan con el recuadro

# clic en la columna PPT del primer recuadro
bb = a.t_rec.bbox(a.t_rec.get_children()[0], "#1")


class Ev:
    x, y = bb[0] + 4, bb[1] + 4


a._rec_click_casilla(Ev())
assert a.recuadros[0]["en_pptx"] is False
assert a.t_rec.raw_values(a.t_rec.get_children()[0])[0] == CASILLA_NO
a._rec_marcar_todas(True)
assert all(r.get("en_pptx", True) for r in a.recuadros)
assert a.t_rec.raw_values(a.t_rec.get_children()[0])[0] == CASILLA_SI

# editar un recuadro y cambiar su tipo no pierde lo escrito en campos comunes
a.t_rec.selection_set(a.t_rec.get_children()[1])
a._rec_editar()
a.rv_tipo.set("Homicidio / Lesiones graves")
a._rec_cambiar_tipo()
assert "Lugar" in a._rec_widgets
assert a._rec_widgets["Lugar"][0].get() == "San Antonio"
a._rec_guardar()
assert not errores, errores

# ── 7. Texto X ───────────────────────────────────────────────────────────────
a._switch("textox")
a.txt_x.insert("1.0", "Resumen de relevancia.")
a._doc_cambio()

# ── 8. Fotos de las láminas y exportación de punta a punta ───────────────────
a._fotos_dialogo()
a.update()
top = [x for x in a.winfo_children() if x.winfo_class() == "Toplevel"][0]
top.destroy()
a.fotos_pptx["armas_1"] = foto
a.fotos_pptx["det_2"] = os.path.join(tmp, "no_existe.jpg")   # foto perdida: aviso, no falla
salida = os.path.join(tmp, "informe.pptx")
filedialog.asksaveasfilename = lambda **k: salida
cuadros.clear()
a._exportar_pptx()
a.update()
assert os.path.isfile(salida), "no se generó el PowerPoint"
prs = Presentation(salida)
assert len(prs.slides) == 8 + 2 + 1, len(prs.slides)
fotos_laminas = {i for i, s in enumerate(prs.slides, 1)
                 for sh in s.shapes if sh.shape_type == 13}
assert 2 in fotos_laminas, fotos_laminas              # la foto de armas quedó en la lámina 2
assert not [x for x in cuadros if x[0] == "error"], cuadros

# casilla PPT desmarcada: ese recuadro no sale
a.recuadros[0]["en_pptx"] = False
a._exportar_pptx()
assert len(Presentation(salida).slides) == 8 + 1 + 1
assert not os.path.exists(salida + ".tmp")
assert not [f for f in os.listdir(tmp) if f.endswith(".pptx") and f != "informe.pptx"]

# exportar con el archivo bloqueado no deja basura ni cierra la app
import comite_policial.exporters.pptx_export as PX
original = PX.exportar_informe
PX.exportar_informe = lambda *x, **k: (_ for _ in ()).throw(PermissionError("abierto"))
cuadros.clear()
a._exportar_pptx()
assert any(k == "warn" for k, _ in cuadros), cuadros
PX.exportar_informe = original

# plantilla ausente
from comite_policial.config import settings as S
import dataclasses
PX.settings = dataclasses.replace(S.settings, plantilla_pptx=os.path.join(tmp, "nada.pptx"))
cuadros.clear()
a._exportar_pptx()
assert any(k == "error" for k, _ in cuadros), cuadros
PX.settings = S.settings

# ── 9. Importar un Word: éxito, archivo dañado y archivo ajeno ───────────────
docx_ok = make_informe_docx.crear(os.path.join(tmp, "informe.docx"))
filedialog.askopenfilename = lambda **k: docx_ok
cuadros.clear()
a._doc_importar()
a.update()
assert a._doc is not None and len(a.recuadros) == 8, len(a.recuadros)
assert a.v_semana.get() == "31.AGO AL 06.SEP.026", a.v_semana.get()
assert a.t_armas.raw_values(a.t_armas.get_children()[0])[1] == "6"
assert a._nav["armas"]._partes[3].cget("text") == "●"  # el menú marca la sección con datos
assert not [x for x in cuadros if x[0] == "error"], cuadros

malo = os.path.join(tmp, "malo.docx")
open(malo, "wb").write(b"esto no es un docx")
filedialog.askopenfilename = lambda **k: malo
cuadros.clear()
doc_previo = a._doc
a._doc_importar()
assert any(k == "error" for k, _ in cuadros), cuadros
assert a._doc is doc_previo, "un archivo dañado no debe reemplazar el documento abierto"

ajeno = os.path.join(tmp, "ajeno.docx")
from docx import Document
d = Document(); d.add_paragraph("Un documento cualquiera"); d.save(ajeno)
filedialog.askopenfilename = lambda **k: ajeno
cuadros.clear()
a._doc_importar()
assert any(k == "warn" for k, _ in cuadros), cuadros
assert a._doc is doc_previo

# ── 10. Sincronización con el Word y cierre ordenado ─────────────────────────
t = a.t_armas
fila = t.get_children()[0]
v = t.raw_values(fila)
v[1] = "11"
t.item(fila, values=v)
a._recalc_armas()
a._doc_cambio()                                        # queda pendiente (700 ms)
assert a._doc_after is not None
a._cerrar()                                            # al cerrar se escribe lo pendiente
r = importar(docx_ok)
assert r["tablas"]["armas"]["rows"][0][1] == "11", r["tablas"]["armas"]["rows"][0]

assert not errores, errores
print("OK")
